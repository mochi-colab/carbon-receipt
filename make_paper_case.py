"""공개 논문 1편의 보고값을 입력양식에 옮겨 엔진을 실행한다.

논문: Al-Fatesh, A.S. et al. (2016) Effect of Ce and Co Addition to Fe/Al2O3 for Catalytic Methane
Decomposition. Catalysts 6(3), 40. https://doi.org/10.3390/catal6030040 (CC BY 3.0)
대상 촉매: 30%Fe/Al2O3 (500 ℃ 환원, 700 ℃ 반응). 쪽 번호는 출판 PDF(15쪽) 기준.

원칙: 논문에 보고된 값만 넣는다. 보고되지 않은 값은 비워 두고 이유를 비고에 적는다.
1차 실행: 보고값만 / 2차 실행: 전력만 가정값([가정값] 표시)을 추가.
"""
import csv
import os
import shutil

from openpyxl import load_workbook

import carbon_receipt as cr
from make_cases import FORM, example_rows, write_case

OUT = os.path.join("결과", "논문적용")
EXP = "LIT-AF2016-Fe30"
REF = "Al-Fatesh et al. (2016) Catalysts 6, 40"

# (구간, 항목코드) → (값, 단위, 비고). 없는 키는 빈칸 + 기본 사유.
REPORTED = {
    ("반응", "TEMP"): (700, "℃", f"보고값. {REF} p.13(3.3절) 'desired reaction temperature (i.e., 700 °C)'"),
    ("반응", "OP_TIME"): (180, "min", f"보고값. {REF} p.6(2.2.1절) 'After 180 min', p.10 그림 12 캡션"),
    ("반응", "CAT_TYPE"): ("Fe", "-", f"보고값. {REF} p.11(3.1절) 30 wt% Fe/Al2O3, 습식 함침·450 ℃ 소성. "
                                       "배출계수는 Fe 예시계수에 연결(담체 Al2O3는 계수에 반영 안 됨)"),
    ("반응", "CAT_MASS"): (0.3, "g", f"보고값. {REF} p.13(3.3절) 'a fixed mass of catalyst (0.3 g)'"),
}
NOT_REPORTED = {
    ("승온", "OP_TIME"): "미보고. 환원(H2 40 mL/min, 90 min, p.13)은 기재되나 승온 시간은 없음",
    ("승온", "ELEC"): "미보고. 논문 전체에 전력·에너지 사용량 없음",
    ("승온", "HEAT_EXT"): "미보고. 반응기 가열 방식 기재 없음(p.13)",
    ("반응", "CH4_NEW"): "미보고. 재순환 여부 기재 없음",
    ("반응", "CH4_IN"): "질량 미보고. p.13 'methane/nitrogen was 1.5/1', 'total flow rate was 25 mL/min'와 반응 180 min에서 "
                        "CH4 15 mL/min × 180 min = 2.7 L(체적)로 계산되나, 질량 환산에 필요한 기준 온도·압력이 없어 비움",
    ("반응", "CH4_OUT"): "미보고. 메탄 전환율 수치 없음(p.2에는 목표 문장만)",
    ("반응", "X_CH4"): "미보고. 출구 가스 조성 수치 없음",
    ("반응", "X_H2"): "미보고. 출구 가스 조성 수치 없음",
    ("반응", "ELEC"): "미보고. 논문 전체에 전력·에너지 사용량 없음",
    ("반응", "HEAT_EXT"): "미보고. 반응기 가열 방식 기재 없음(p.13)",
    ("반응", "H2_PROD"): "질량 미보고. p.6(2.2.1절)·그림 5에 H2 수율 54.9%(20 min)·69%(180 min)만 있고 수율 정의식이 없어 "
                         "질량으로 환산하지 않음",
    ("반응", "H2_PURITY"): "미보고",
    ("반응", "C_REC"): "미보고. 탄소는 TPO·TEM 정성 결과만 있음(p.10–11)",
    ("반응", "C_ACC"): "미보고",
    ("반응", "CAT_USES"): "미보고. 촉매 재사용·재생 여부 기재 없음",
    ("중단", "OP_TIME"): "미보고",
    ("중단", "ELEC"): "미보고. 논문 전체에 전력·에너지 사용량 없음",
    ("중단", "HEAT_EXT"): "미보고. 반응기 가열 방식 기재 없음(p.13)",
}
# 2차 실행에만 넣는 전력 가정값: 출처 = 본 보고서 작성예시(EXP-800)의 같은 구간 값. 실제 장비 전력 아님.
ASSUMED_ELEC = {"승온": 1200, "반응": 5000, "중단": 300}


def rows_for(with_assumed_power):
    wb = load_workbook(FORM)
    rows = []
    for r in example_rows(wb):
        seg, code = r[1], r[3]
        new = list(r)
        new[0] = EXP
        if (seg, code) in REPORTED:
            v, u, note = REPORTED[(seg, code)]
            new[4], new[5], new[9] = v, u, note
        else:
            new[4] = None
            new[9] = NOT_REPORTED[(seg, code)]
            if with_assumed_power and code == "ELEC":
                new[4], new[5] = ASSUMED_ELEC[seg], "Wh"
                new[9] = ("[가정값] 논문 미보고. 시연용으로 본 보고서 작성예시(EXP-800)의 같은 구간 값을 그대로 넣음 "
                          "— 실제 장비 전력 아님")
        rows.append(new)
    return rows


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT)
    results = {}
    for tag, assumed in (("1차_보고값만", False), ("2차_전력가정값", True)):
        path = os.path.join(OUT, f"논문입력_{tag}.xlsx")
        write_case(rows_for(assumed), path)
        # 자료 상태 표시 문구를 이 파일에 맞게 바꿈(가상 데이터가 아니라 논문 보고값)
        wb = load_workbook(path)
        ws = wb["계산설정"]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == "LABEL_DATA":
                ws.cell(r, 4).value = ("공개 논문 보고값(Al-Fatesh 2016) · 예시 배출계수 기반" +
                                       (" · 전력은 가정값" if assumed else "") + " (공인 계수 아님)")
        wb.save(path)
        res = cr.run(path, os.path.join(OUT, f"영수증_{tag}"))[0]
        results[tag] = res
    # 항목 대조표: 논문에 보고된 항목 / 보고되지 않은 항목
    with open(os.path.join(OUT, "논문_항목대조.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["구간", "항목", "항목코드", "필수/선택", "논문 보고", "입력값", "단위", "출처·사유"])
        for r in rows_for(False):
            rep = (r[1], r[3]) in REPORTED
            w.writerow([r[1], r[2], r[3], r[7], "보고됨" if rep else "보고 안 됨",
                        "" if r[4] is None else r[4], r[5] if rep else "", r[9]])
    for tag, res in results.items():
        print(f"== {tag}: {res['state']} ({res['state_reason']})")
        for k, (s, d) in res["checks"].items():
            print(f"   {k:<16} {s:<8} {d}")
        for m in res["messages"]:
            if m.startswith("[자료 부족]"):
                print("   ·", m[:90])
        if res["assumed"]:
            print("   가정값:", "; ".join(res["assumed"]))
    return results


if __name__ == "__main__":
    main()
