# 탄소영수증 입력양식 생성 스크립트 (Prompt A)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
WARN_FILL = PatternFill("solid", fgColor="FDE9E7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)

def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = f(True, "FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

def body(cell, fill=None, wrap=True):
    cell.font = f()
    cell.border = BOX
    cell.alignment = Alignment(vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = fill

G, A, L = "기록될 것으로 예상", "추가 측정이 필요할 것으로 예상", "확인 전 판단 어려움"

# (구간, 항목, 항목코드, 기본단위, 측정 구간 설명, 필수/선택, 확보 구분, 비고, 예시값, 예시단위)
ITEMS = [
    ("승온", "운전시간", "OP_TIME", "h", "가열 시작 ~ 반응온도 도달(메탄 공급 시작 전)", "선택", L,
     "승온 프로그램은 설정되지만 구간별 실제 소요시간을 따로 남기는지 확인 필요", 1.5, "h"),
    ("승온", "전력 사용량", "ELEC", "kWh", "가열 시작 ~ 메탄 공급 시작, 반응로·보조장비 전체", "필수", A,
     "실험실 전기로는 전력량을 기록하지 않는 경우가 많아 전력량계 설치·로그 확인 필요", 1200, "Wh"),
    ("승온", "외부 공급열", "HEAT_EXT", "MJ", "승온 구간에 전기 외 열원(버너·예열기 등)으로 공급한 열", "필수", L,
     "전기 가열만 쓰면 0 입력. 다른 열원 사용 여부는 장비 구성에 따라 다름", 0, "MJ"),
    ("반응", "메탄 신규 공급량", "CH4_NEW", "kg", "메탄 공급 시작 ~ 중단, 외부에서 새로 공급한 메탄(재순환분 제외)", "필수", L,
     "재순환 없는 일회 통과식이면 투입량과 같음. 재순환 구성·집계 방식 확인 필요", 1.0, "kg"),
    ("반응", "반응기 메탄 투입량", "CH4_IN", "kg", "메탄 공급 시작 ~ 중단, 반응기 입구 기준(재순환분 포함)", "선택", G,
     "MFC 유량 설정값×공급시간으로 환산해 기록하는 경우가 많음. 화학양론 점검용", 2.0, "kg"),
    ("반응", "반응기 메탄 배출량(미반응)", "CH4_OUT", "kg", "메탄 공급 시작 ~ 중단, 반응기 출구의 미반응 메탄", "선택", L,
     "출구 유량 실측인지, GC 조성·전환율로 역산하는지 연구실마다 다름. 화학양론 점검용", 1.0, "kg"),
    ("반응", "가스 조성 - CH4 몰분율", "X_CH4", "mol/mol", "반응 구간 출구 가스, GC 분석(평균값)", "선택", G,
     "메탄 열분해 실험에서 GC 출구 조성 분석은 통상 수행됨", 0.334, "mol/mol"),
    ("반응", "가스 조성 - H2 몰분율", "X_H2", "mol/mol", "반응 구간 출구 가스, GC 분석(평균값)", "선택", G,
     "메탄 열분해 실험에서 GC 출구 조성 분석은 통상 수행됨", 0.666, "mol/mol"),
    ("반응", "반응온도", "TEMP", "℃", "반응 구간 설정(유지) 온도", "필수", G,
     "실험조건으로 항상 기록됨. 영수증 식별용(계산에는 미사용)", 800, "℃"),
    ("반응", "운전시간", "OP_TIME", "h", "메탄 공급 시작 ~ 중단", "필수", G,
     "반응시간(TOS)은 실험조건으로 통상 기록됨", 4.0, "h"),
    ("반응", "전력 사용량", "ELEC", "kWh", "메탄 공급 시작 ~ 중단, 반응로·보조장비 전체", "필수", A,
     "실험실 전기로는 전력량을 기록하지 않는 경우가 많아 전력량계 설치·로그 확인 필요", 5000, "Wh"),
    ("반응", "외부 공급열", "HEAT_EXT", "MJ", "반응 구간에 전기 외 열원으로 공급한 열", "필수", L,
     "전기 가열만 쓰면 0 입력. 다른 열원 사용 여부는 장비 구성에 따라 다름", 0, "MJ"),
    ("반응", "수소 생산량", "H2_PROD", "kg", "메탄 공급 시작 ~ 중단, 출구 가스 중 H2 질량(유량×H2 몰분율 적산)", "필수", L,
     "논문·노트에는 전환율·수율로 남기는 경우가 많아 질량 환산 가능 여부 확인 필요", 250, "g"),
    ("반응", "수소 순도", "H2_PURITY", "mol%", "수소 생산량을 측정한 지점의 H2 몰분율(정제 전)", "선택", L,
     "'순도'를 출구 가스 기준으로 볼지 분리 후 기준으로 볼지 정의 확인 필요. 정제는 경계 밖", 66.6, "mol%"),
    ("반응", "고체탄소 회수량", "C_REC", "kg", "실험 종료 후 회수·칭량한 고체탄소(촉매 질량 제외)", "선택", L,
     "촉매 전후 무게 차로 탄소 수율을 보고하는 경우가 있으나 회수·칭량 방식 확인 필요. 질량 배분·탄소 점검용", 720, "g"),
    ("반응", "반응기 내 탄소 축적량", "C_ACC", "kg", "회수하지 못하고 반응기 벽·배관에 남은 탄소", "선택", A,
     "통상 별도로 측정하지 않음. 탄소수지 점검 정확도를 높일 때 추가 측정", 20, "g"),
    ("반응", "촉매 종류", "CAT_TYPE", "-", "이번 실험에 사용한 촉매(배출계수 선택에 사용)", "필수", G,
     "실험조건으로 항상 기록됨", "Fe", "-"),
    ("반응", "촉매 사용량", "CAT_MASS", "kg", "반응기에 충전한 촉매 질량(최초 투입 기준)", "필수", G,
     "실험조건으로 통상 기록됨", 20, "g"),
    ("반응", "촉매 최초 사용 포함 총 사용 횟수", "CAT_USES", "회", "같은 촉매 충전분으로 수행한(할) 실험 횟수. 최초 사용을 1회로 셈(예: 최초 1회 + 재사용 3회 = 4)", "필수", L,
     "재생·재사용 여부와 횟수를 실험별로 남기는지 확인 필요. 1이면 이번 실험에 전량 배분", 4, "회"),
    ("중단", "운전시간", "OP_TIME", "h", "메탄 공급 중단 ~ 전원 차단(냉각·퍼지 포함)", "선택", L,
     "냉각·퍼지 시간을 구간별로 남기는지 확인 필요", 1.0, "h"),
    ("중단", "전력 사용량", "ELEC", "kWh", "메탄 공급 중단 ~ 전원 차단, 반응로·보조장비 전체", "필수", A,
     "실험실 전기로는 전력량을 기록하지 않는 경우가 많아 전력량계 설치·로그 확인 필요", 300, "Wh"),
    ("중단", "외부 공급열", "HEAT_EXT", "MJ", "중단 구간에 전기 외 열원으로 공급한 열", "필수", L,
     "전기 가열만 쓰면 0 입력. 다른 열원 사용 여부는 장비 구성에 따라 다름", 0, "MJ"),
]

COLS = ["실험ID", "구간", "항목", "항목코드", "값", "단위", "측정 구간 설명", "필수/선택", "확보 구분", "비고"]
WIDTHS = [11, 7, 24, 11, 10, 9, 44, 9, 16, 58]

def add_validations(ws, n):
    last = n + 1
    for col, opts in (("B", "승온,반응,중단"), ("H", "필수,선택"), ("I", f"{G},{A},{L}")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last + 20}")

def legend(ws, start_col, lines, title):
    c = ws.cell(row=1, column=start_col, value=title)
    c.font = f(True, "FFFFFF"); c.fill = PatternFill("solid", fgColor="7F7F7F")
    ws.column_dimensions[c.column_letter].width = 62
    for i, t in enumerate(lines, start=2):
        cell = ws.cell(row=i, column=start_col, value=t)
        cell.font = f(); cell.fill = NOTE_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---------------- 1. 실험기록 (빈 양식) ----------------
ws = wb.active
ws.title = "실험기록"
header(ws, COLS, WIDTHS)
for r, it in enumerate(ITEMS, start=2):
    seg, name, code, unit, desc, req, src, note, _, _ = it
    vals = ["", seg, name, code, "", unit, desc, req, src, note]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        body(cell, INPUT_FILL if c in (1, 5, 6) else None)
add_validations(ws, len(ITEMS))
ws.column_dimensions["K"].width = 3
legend(ws, 12, [
    "노란 칸(실험ID·값·단위)만 입력한다. 나머지 열은 양식 정의이므로 수정하지 않는다.",
    "실험 1건 = 이 표 전체(승온·반응·중단 구간별 행). 실험이 여러 건이면 행 묶음을 복사해 아래에 붙이고 실험ID를 바꾼다.",
    "단위는 측정 장비에 기록된 그대로 적는다(예: Wh, g). 계산 시 '계산설정'의 단위환산표로 표준단위로 바꾼다.",
    "측정하지 않은 값은 비워 둔다. 추정값으로 채우지 않는다(필수값이 비면 계산 보류, 선택값이 비면 해당 점검·참고값만 '자료 부족').",
    "전기 외 열원이 없으면 외부 공급열은 0을 입력한다(빈칸과 0을 구분).",
    "같은 실험ID·구간·항목 행이 두 번 있으면 중복 기록으로 보고 계산을 보류한다.",
    "작성 예시는 '작성예시' 시트 참고(가상 데이터).",
], "작성 안내")

# ---------------- 2. 계산설정 ----------------
ws = wb.create_sheet("계산설정")
header(ws, ["설정ID", "구분", "항목", "값", "단위", "설명·근거"], [18, 12, 30, 44, 12, 60])
S = [
    ("FORM_VERSION", "버전", "입력양식 버전", "v0.1", "-", "양식 구조 변경 시 갱신"),
    ("EF_VERSION", "버전", "적용 배출계수 버전", "EF-DEMO-v0.1", "-", "'배출계수' 시트의 버전 열과 일치하는 계수만 사용"),
    ("FU", "기능단위", "기능단위", 1, "kg-H2", "수소 1 kg당 배출량(kgCO2-eq/kg-H2)으로 표시"),
    ("BOUNDARY_IN_1", "시스템 경계", "포함", "원료 메탄 생산·공급", "-", "메탄 신규 공급량 × 메탄 생산·공급 배출계수"),
    ("BOUNDARY_IN_2", "시스템 경계", "포함", "공정 전력·열", "-", "전력 사용량, 외부 공급열"),
    ("BOUNDARY_IN_3", "시스템 경계", "포함", "촉매", "-", "촉매 제조 배출(실험 1회 배분량 기준)"),
    ("BOUNDARY_OUT_1", "시스템 경계", "제외", "수소 정제·압축", "-", "실험실 규모에서 정제·압축 공정 없음. 영수증에 경계 밖 항목으로 표시"),
    ("BOUNDARY_OUT_2", "시스템 경계", "제외", "공정 내 메탄 누출", "-", "실험실 배관 누출 측정 자료 없음. 영수증에 경계 밖 항목으로 표시"),
    ("BOUNDARY_OUT_3", "시스템 경계", "제외", "고체탄소 후처리·활용", "-", "회수 이후 정제·판매·대체효과는 본 단계에서 다루지 않음"),
    ("BOUNDARY_OUT_4", "시스템 경계", "제외", "미반응 메탄의 배기·연소 처리", "-",
     "예시는 미반응 메탄 회수·재순환을 가정. 일회 통과식 반응기라면 CH4_NEW=CH4_IN으로 입력하고, 미반응 메탄 처리 배출의 포함 여부는 연구실 확인 후 결정"),
    ("LABEL_RESULT", "결과 표시", "결과 표시 문구", "실험실 규모 연구용 결과 — 정제·압축 제외, 상업 규모 추정치 아님", "-", "모든 영수증에 표시"),
    ("LABEL_DATA", "결과 표시", "자료 상태 표시 문구", "가상 데이터·예시 배출계수 기반 (실제 실험값·공인 계수 아님)", "-", "실제 자료·확인된 계수로 바꾸기 전까지 표시"),
    ("ENERGY_SCOPE_BASE", "에너지 집계", "기본 집계 구간", "승온+반응+중단", "-", "승온·반응·중단 구간 전력을 합산해 같은 배치의 수소 생산량으로 나눔. 배치의 시작·종료 시점과 냉각·대기 구간 포함 여부는 실제 적용 전 연구실과 합의"),
    ("ENERGY_SCOPE_REF", "에너지 집계", "참고 집계 구간", "반응", "-", "반응 구간만 별도 계산(참고값). 승온·중단 비중을 확인하는 용도"),
    ("CH4_BASIS", "메탄 가정", "원료 배출 산정 기준", "CH4_NEW", "-", "메탄 신규 공급량으로 계산. 미반응 메탄 회수·재순환은 신규 공급량에 반영된 것으로 봄"),
    ("CAT_ALLOC", "촉매 가정", "실험 1회 배분량", "CAT_MASS / CAT_USES", "kg", "촉매 사용량 ÷ 최초 사용 포함 총 사용 횟수"),
    ("CARBON_BASE", "고체탄소 가정", "기본", "부산물 공제 없음", "-", "공정 부담 전량을 수소에 귀속"),
    ("CARBON_REF", "고체탄소 가정", "참고", "질량 배분", "-", "고체탄소 회수량(C_REC)이 있을 때만 계산. 수소 배분율 = H2 ÷ (H2 + C_REC)"),
    ("C_BALANCE_BASIS", "검사 기준", "탄소 점검 비교값", "C_REC + C_ACC", "kg", "고체탄소 회수량 + 반응기 내 탄소 축적량(기록된 경우)을 이론 탄소량과 비교"),
    ("CALC_TOL", "검사 기준", "계산 허용 오차(수작업 대비)", 0.001, "상대오차", "자동 계산과 수작업 기준 계산의 상대오차 0.1% 이내"),
    ("STOICH_TOL", "검사 기준", "화학양론 일관성 허용 범위", 0.10, "상대오차", "예시값. 반응 메탄 기준 H2·C 이론량 대비 ±10% 밖이면 '확인 필요'. 실제 적용 시 측정 오차 수준에 맞춰 조정"),
    ("ELEC_INTENSITY_MIN", "검사 기준", "전력 원단위 하한", 5, "kWh/kg-H2", "예시값. 배치 전체 전력 ÷ 수소 생산량이 범위 밖이면 단위 오기 의심 '확인 필요'"),
    ("ELEC_INTENSITY_MAX", "검사 기준", "전력 원단위 상한", 100, "kWh/kg-H2", "예시값. 실제 적용 시 동일 장비 과거 실험 범위로 교체"),
    ("RULE_MISSING", "검사 규칙", "필수값 누락", "자료 부족 · 계산 보류", "-", "필요한 항목명과 측정 구간을 안내"),
    ("RULE_DUPLICATE", "검사 규칙", "중복 기록", "확인 필요 · 계산 보류", "-", "같은 실험ID·구간·항목이 2행 이상. 자동 삭제하지 않음"),
    ("RULE_H2_NONPOSITIVE", "검사 규칙", "수소 생산량 0 이하", "계산 불가", "-", "기능단위 환산 불가(0으로 나누기 방지). 사유 표시"),
    ("RULE_ELEC_OUTLIER", "검사 규칙", "전력 원단위 범위 이탈", "확인 필요 · 계산 보류", "-", "단위 오기 의심. 원자료 대조·수정 후 재검사"),
    ("ACK_ALLOWED", "검사 규칙", "연구자 확인 후 계산 진행이 허용되는 검사", "전력 원단위 이상치", "-", "원자료 대조로 값이 맞다고 확인된 경우에만 사용. 경고는 영수증에 그대로 남김. 필수값 누락·중복 기록·단위 미인식은 확인 표시로 넘어갈 수 없음"),
    ("RULE_RANGE", "검사 규칙", "값 범위·물리 한계 위반", "확인 필요 · 계산 보류", "-", "음수 소비량, 몰분율·온도 범위 밖, 정수가 아닌 사용 횟수, 투입보다 많은 미반응 메탄, 수소 생산량에 비해 부족한 메탄 신규 공급량, 등록되지 않은 구간·항목코드"),
    ("RULE_EF_INVALID", "검사 규칙", "배출계수 단위·값·중복 문제", "확인 필요 · 계산 보류", "-", "계수 단위의 분모가 활동량 표준단위와 다름(자동 환산하지 않음), 음수 계수, 같은 버전·항목코드에 계수 2개 이상"),
    ("REQUIRED_ROWS", "검사 기준", "필수 행(구간:항목코드)", "승온:ELEC, 승온:HEAT_EXT, 반응:CH4_NEW, 반응:TEMP, 반응:OP_TIME, 반응:ELEC, 반응:HEAT_EXT, 반응:H2_PROD, 반응:CAT_TYPE, 반응:CAT_MASS, 반응:CAT_USES, 중단:ELEC, 중단:HEAT_EXT", "-", "입력 파일의 필수/선택 표시와 별도로, 이 목록의 행이 없거나 값이 비어 있으면 자료 부족·계산 보류"),
    ("ALLOWED_SEGMENTS", "검사 기준", "허용 구간", "승온, 반응, 중단", "-", "다른 이름의 구간 행은 집계에서 빠지므로 '확인 필요'"),
    ("NONNEG_CODES", "검사 기준", "음수가 될 수 없는 항목코드", "ELEC, HEAT_EXT, CH4_NEW, CH4_IN, CH4_OUT, H2_PROD, C_REC, C_ACC, CAT_MASS, OP_TIME", "-", "질량·에너지·시간 항목. 0은 허용"),
    ("INT_MIN1_CODES", "검사 기준", "1 이상의 정수여야 하는 항목코드", "CAT_USES", "-", "촉매 최초 사용 포함 총 사용 횟수"),
    ("POSITIVE_ROWS", "검사 기준", "0보다 커야 하는 행(구간:항목코드)", "반응:OP_TIME", "-", "수소가 생산된 반응 구간의 운전시간이 0이면 기록 오류로 봄"),
    ("SUM_MAX_X", "검사 기준", "몰분율 합 상한(X_CH4 + X_H2)", 1, "mol/mol", "출구 가스 성분 몰분율의 합은 1을 넘을 수 없음"),
    ("RANGE_X_CH4", "검사 기준", "X_CH4", "0~1", "mol/mol", "몰분율 허용 범위"),
    ("RANGE_X_H2", "검사 기준", "X_H2", "0~1", "mol/mol", "몰분율 허용 범위"),
    ("RANGE_H2_PURITY", "검사 기준", "H2_PURITY", "0~1", "mol/mol", "표준단위(mol/mol) 환산 후 허용 범위"),
    ("RANGE_TEMP", "검사 기준", "TEMP", "0~1500", "℃", "예시값. 메탄 열분해 실험의 반응온도로 볼 수 있는 범위"),
    ("CH4_BALANCE_RULE", "검사 기준", "메탄 질량 관계", "CH4_OUT ≤ CH4_IN; CH4_NEW ≥ H2_PROD×M_CH4÷(2×M_H2)×(1−STOICH_TOL)", "-", "미반응 메탄이 투입량보다 많거나, 메탄 신규 공급량이 보고된 수소를 만들 수 있는 이론 최소량보다 적으면 '확인 필요'"),
    ("EFU_ELEC", "배출계수 단위", "ELEC", "kgCO2-eq/kWh", "-", "활동량 표준단위(kWh)에 맞는 계수 단위. 다르면 자동 환산하지 않고 보류"),
    ("EFU_HEAT_EXT", "배출계수 단위", "HEAT_EXT", "kgCO2-eq/MJ", "-", ""),
    ("EFU_CH4_NEW", "배출계수 단위", "CH4_NEW", "kgCO2-eq/kg-CH4", "-", ""),
    ("EFU_CAT_MASS", "배출계수 단위", "CAT_MASS", "kgCO2-eq/kg-촉매", "-", ""),
    ("RULE_STOICH", "검사 규칙", "화학양론 불일치", "확인 필요 · 계산 수행(경고 표시)", "-", "가능 원인(입력 오류 / 반응기 내 축적 / 미측정 유출)을 영수증에 표시"),
    ("RULE_UNIT_UNKNOWN", "검사 규칙", "인식할 수 없는 단위", "확인 필요 · 계산 보류", "-", "단위환산표에 없거나 표준단위와 차원이 다른 단위"),
    ("RULE_EF_MISSING", "검사 규칙", "배출계수 없음", "자료 부족 · 계산 보류", "-", "적용 버전의 배출계수 중 항목코드에 맞는 계수가 없음(예: 촉매 종류)"),
    ("STD_OP_TIME", "표준단위", "OP_TIME", "h", "-", "항목코드별 표준단위. 입력단위가 이와 다르면 단위환산표로 환산"),
    ("STD_ELEC", "표준단위", "ELEC", "kWh", "-", ""),
    ("STD_HEAT_EXT", "표준단위", "HEAT_EXT", "MJ", "-", ""),
    ("STD_CH4_NEW", "표준단위", "CH4_NEW", "kg", "-", ""),
    ("STD_CH4_IN", "표준단위", "CH4_IN", "kg", "-", ""),
    ("STD_CH4_OUT", "표준단위", "CH4_OUT", "kg", "-", ""),
    ("STD_X_CH4", "표준단위", "X_CH4", "mol/mol", "-", ""),
    ("STD_X_H2", "표준단위", "X_H2", "mol/mol", "-", ""),
    ("STD_TEMP", "표준단위", "TEMP", "℃", "-", ""),
    ("STD_H2_PROD", "표준단위", "H2_PROD", "kg", "-", ""),
    ("STD_H2_PURITY", "표준단위", "H2_PURITY", "mol/mol", "-", ""),
    ("STD_C_REC", "표준단위", "C_REC", "kg", "-", ""),
    ("STD_C_ACC", "표준단위", "C_ACC", "kg", "-", ""),
    ("STD_CAT_TYPE", "표준단위", "CAT_TYPE", "-", "-", ""),
    ("STD_CAT_MASS", "표준단위", "CAT_MASS", "kg", "-", ""),
    ("STD_CAT_USES", "표준단위", "CAT_USES", "회", "-", ""),
    ("M_CH4", "물성", "CH4 몰질량", 16.043, "g/mol", "화학양론 이론량 계산용"),
    ("M_H2", "물성", "H2 몰질량", 2.016, "g/mol", "화학양론 이론량 계산용"),
    ("M_C", "물성", "C 몰질량", 12.011, "g/mol", "화학양론 이론량 계산용"),
    ("UNIT_Wh", "단위환산", "Wh", 0.001, "kWh", "입력단위 → 표준단위 곱하는 값"),
    ("UNIT_kWh", "단위환산", "kWh", 1, "kWh", ""),
    ("UNIT_MWh", "단위환산", "MWh", 1000, "kWh", ""),
    ("UNIT_g", "단위환산", "g", 0.001, "kg", ""),
    ("UNIT_kg", "단위환산", "kg", 1, "kg", ""),
    ("UNIT_mg", "단위환산", "mg", 0.000001, "kg", ""),
    ("UNIT_kJ", "단위환산", "kJ", 0.001, "MJ", ""),
    ("UNIT_MJ", "단위환산", "MJ", 1, "MJ", ""),
    ("UNIT_min", "단위환산", "min", 1 / 60, "h", ""),
    ("UNIT_h", "단위환산", "h", 1, "h", ""),
    ("UNIT_molpct", "단위환산", "mol%", 0.01, "mol/mol", "몰분율 입력 시"),
    ("UNIT_molmol", "단위환산", "mol/mol", 1, "mol/mol", ""),
]
for r, row in enumerate(S, start=2):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        body(cell)
    if "예시값" in row[5]:
        ws.cell(row=r, column=4).fill = WARN_FILL
    if row[0] == "STOICH_TOL" or row[0] == "CALC_TOL":
        ws.cell(row=r, column=4).number_format = "0.0%"
legend(ws, 8, [
    "이 시트의 값은 계산 엔진이 그대로 읽는다. 설정ID는 바꾸지 않는다.",
    "분홍 칸은 시연용 예시값이다. 실제 적용 전에 연구실 자료로 확인·교체한다.",
    "에너지 집계 구간·메탄 기준·촉매 배분·고체탄소 처리 가정은 영수증의 '적용 가정'에 그대로 표시된다.",
], "작성 안내")
ws.column_dimensions["G"].width = 3

# ---------------- 3. 배출계수 ----------------
ws = wb.create_sheet("배출계수")
header(ws, ["계수ID", "대상", "값", "단위", "출처", "기준연도", "적용 조건", "버전", "적용 항목코드"],
       [14, 18, 8, 18, 58, 12, 40, 14, 16])
EF = [
    ("EF-ELEC-01", "전력(계통)", 0.45, "kgCO2-eq/kWh",
     "예시값(시연용, 출처 아님). 실제 적용 시 확인할 출처 종류: 국가 고시 전력 배출계수 또는 LCA DB의 국가 전력 믹스 공정",
     "미정(예시)", "연구실 계통전력 사용. 승온·반응·중단 전 구간 전력에 적용", "EF-DEMO-v0.1", "ELEC"),
    ("EF-HEAT-01", "외부 공급열", 0.07, "kgCO2-eq/MJ",
     "예시값(시연용, 출처 아님). 실제 적용 시 확인할 출처 종류: 열원 연료의 연소·상류 배출계수(국가 고유계수 또는 LCA DB)",
     "미정(예시)", "외부 공급열 > 0일 때만 적용. 열원 종류별로 계수 분리 필요", "EF-DEMO-v0.1", "HEAT_EXT"),
    ("EF-CH4-01", "메탄 생산·공급", 0.6, "kgCO2-eq/kg-CH4",
     "예시값(시연용, 출처 아님). 실제 적용 시 확인할 출처 종류: LCA DB의 천연가스(배관/LNG) 공급 공정 또는 실험용 고순도 메탄 제조 공정",
     "미정(예시)", "공급 경로(배관·LNG·고압용기) 가정을 명시해야 함", "EF-DEMO-v0.1", "CH4_NEW"),
    ("EF-CAT-FE", "촉매(Fe)", 2.0, "kgCO2-eq/kg-촉매",
     "예시값(시연용, 출처 아님). 실제 적용 시 확인할 출처 종류: LCA DB의 철 또는 해당 촉매 제조 공정(전구체·담체 포함 여부 확인)",
     "미정(예시)", "촉매 종류 = Fe일 때. 실험 1회 배분량에 적용", "EF-DEMO-v0.1", "CAT_MASS:Fe"),
]
for r, row in enumerate(EF, start=2):
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        body(cell)
    ws.cell(row=r, column=3).fill = WARN_FILL
    ws.cell(row=r, column=5).font = f(color="C00000")
ws.column_dimensions["J"].width = 3
legend(ws, 11, [
    "모든 값은 시연용 예시값이며 공식 배출계수가 아니다.",
    "실제 적용 시 출처·기준연도·적용 지역·공급 경로를 확인한 계수로 교체하고 버전을 올린다(예: EF-v1.0).",
    "적용 항목코드: 계산 엔진이 '실험기록'의 항목코드와 연결하는 키. 촉매는 'CAT_MASS:촉매종류' 형식.",
    "계산설정의 EF_VERSION과 버전이 같은 행만 사용한다.",
], "작성 안내")

# ---------------- 4. 작성예시 ----------------
ws = wb.create_sheet("작성예시")
header(ws, COLS, WIDTHS)
row_of = {}
for r, it in enumerate(ITEMS, start=2):
    seg, name, code, unit, desc, req, src, note, ev, eu = it
    vals = ["EXP-800", seg, name, code, ev, eu, desc, req, src, note]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        body(cell, INPUT_FILL if c in (1, 5, 6) else None)
    row_of[(seg, code)] = r
add_validations(ws, len(ITEMS))
ws.column_dimensions["K"].width = 3

# 정합성 검산 패널 (수식, 참고용)
legend(ws, 12, [
    "가상 실험 1건(EXP-800, 800℃, Fe 촉매)의 작성 예시다. 실제 측정값이 아니다.",
    "단위는 일부러 Wh·g·kg를 섞어 적었다(단위 환산 확인용).",
], "작성 안내")
ws.column_dimensions["M"].width = 14
hdr = ws.cell(row=5, column=12, value="정합성 검산(참고, 수식)")
hdr.font = f(True, "FFFFFF"); hdr.fill = PatternFill("solid", fgColor="7F7F7F")
h = ws.cell(row=5, column=13, value="값"); h.font = f(True, "FFFFFF"); h.fill = PatternFill("solid", fgColor="7F7F7F")

def E(seg, code):
    return f"E{row_of[(seg, code)]}"

checks = [
    ("반응 메탄량 (kg) = 투입 − 배출", f"={E('반응','CH4_IN')}-{E('반응','CH4_OUT')}", "0.000"),
    ("H2 이론량 (kg) = 반응 메탄 × 2×M_H2/M_CH4",
     "=M6*2*INDEX(계산설정!D:D,MATCH(\"M_H2\",계산설정!A:A,0))/INDEX(계산설정!D:D,MATCH(\"M_CH4\",계산설정!A:A,0))", "0.0000"),
    ("H2 기록값 (kg)", f"={E('반응','H2_PROD')}/1000", "0.0000"),
    ("H2 기록/이론", "=M8/M7", "0.0%"),
    ("C 이론량 (kg) = 반응 메탄 × M_C/M_CH4",
     "=M6*INDEX(계산설정!D:D,MATCH(\"M_C\",계산설정!A:A,0))/INDEX(계산설정!D:D,MATCH(\"M_CH4\",계산설정!A:A,0))", "0.0000"),
    ("C 기록값 (kg) = 회수 + 축적", f"=({E('반응','C_REC')}+{E('반응','C_ACC')})/1000", "0.0000"),
    ("C 기록/이론", "=M11/M10", "0.0%"),
    ("배치 전체 전력 (kWh)", f"=({E('승온','ELEC')}+{E('반응','ELEC')}+{E('중단','ELEC')})/1000", "0.00"),
    ("전력 원단위 (kWh/kg-H2), 배치 전체", "=M13/M8", "0.0"),
    ("전력 원단위 (kWh/kg-H2), 반응 구간", f"={E('반응','ELEC')}/1000/M8", "0.0"),
    ("촉매 1회 배분량 (kg)", f"={E('반응','CAT_MASS')}/1000/{E('반응','CAT_USES')}", "0.0000"),
    ("신규 메탄 = 반응 메탄 여부", f"=IF(ABS({E('반응','CH4_NEW')}-M6)<1E-9,\"일치(재순환 가정과 정합)\",\"불일치\")", "@"),
    ("출구 몰분율 합(CH4+H2)", f"={E('반응','X_CH4')}+{E('반응','X_H2')}", "0.000"),
]
for i, (lab, frm, fmt) in enumerate(checks, start=6):
    a = ws.cell(row=i, column=12, value=lab); body(a, NOTE_FILL)
    b = ws.cell(row=i, column=13, value=frm); body(b); b.number_format = fmt

wb.save("탄소영수증_입력양식.xlsx")
print("saved", row_of)
