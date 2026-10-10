"""연구자 확인 후 입력값 수정 + 수정 이력 기록

원본 파일은 그대로 두고, 수정한 사본을 저장한다. 사본에는 '수정이력' 시트가 추가(또는 누적)되며,
계산 엔진은 이 시트를 읽어 영수증에 수정 이력을 표시한다.

사용 예:
python correct_record.py 원본.xlsx 수정본.xlsx --exp EXP-800 --seg 반응 --code ELEC \
       --value 5000 --unit Wh --reason "원본 로그 대조: 단위 오기" --by 연구자
"""
import argparse
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

HIST_COLS = ["일시", "실험ID", "구간", "항목코드", "항목", "수정 전 값", "수정 전 단위",
             "수정 후 값", "수정 후 단위", "사유", "수정자"]


def apply_correction(src, dst, exp, seg, code, value, unit, reason, by, when=None, sheet="실험기록"):
    wb = load_workbook(src)
    ws = wb[sheet]
    hdr = {c.value: c.column for c in ws[1] if c.value}
    hits = []
    for r in range(2, ws.max_row + 1):
        if (str(ws.cell(r, hdr["실험ID"]).value).strip() == exp
                and str(ws.cell(r, hdr["구간"]).value).strip() == seg
                and str(ws.cell(r, hdr["항목코드"]).value).strip() == code):
            hits.append(r)
    if len(hits) != 1:
        raise ValueError(f"수정 대상 행이 {len(hits)}개입니다(1개여야 함): {exp}/{seg}/{code}")
    r = hits[0]
    old_v, old_u = ws.cell(r, hdr["값"]).value, ws.cell(r, hdr["단위"]).value
    ws.cell(r, hdr["값"]).value = value
    ws.cell(r, hdr["단위"]).value = unit
    name = ws.cell(r, hdr["항목"]).value

    if "수정이력" in wb.sheetnames:
        hs = wb["수정이력"]
    else:
        hs = wb.create_sheet("수정이력")
        for i, c in enumerate(HIST_COLS, start=1):
            cell = hs.cell(1, i, c)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            hs.column_dimensions[cell.column_letter].width = 16 if c != "사유" else 40
    when = when or datetime.now().strftime("%Y-%m-%d %H:%M")
    hs.append([when, exp, seg, code, name, old_v, old_u, value, unit, reason, by])
    wb.save(dst)
    return {"row": r, "old": (old_v, old_u), "new": (value, unit)}


def main():
    import sys
    if "--ack" in sys.argv:
        ap = argparse.ArgumentParser(description="'확인 필요' 검사에 연구자 확인 표시 남기기")
        ap.add_argument("--ack", action="store_true"); ap.add_argument("src"); ap.add_argument("dst")
        ap.add_argument("--exp", required=True); ap.add_argument("--check", required=True)
        ap.add_argument("--note", required=True); ap.add_argument("--by", default="연구자")
        a = ap.parse_args()
        d = acknowledge_check(a.src, a.dst, a.exp, a.check, a.note, a.by)
        print(f"{a.dst} 저장: '{a.check}' 확인 표시 (확인 당시 판정값: {d})")
        return
    ap = argparse.ArgumentParser(description="입력값 수정 및 수정 이력 기록")
    ap.add_argument("src"); ap.add_argument("dst")
    ap.add_argument("--exp", required=True); ap.add_argument("--seg", required=True)
    ap.add_argument("--code", required=True); ap.add_argument("--value", required=True)
    ap.add_argument("--unit", required=True); ap.add_argument("--reason", required=True)
    ap.add_argument("--by", default="연구자")
    a = ap.parse_args()
    v = a.value
    try:
        v = float(v)
    except ValueError:
        pass
    info = apply_correction(a.src, a.dst, a.exp, a.seg, a.code, v, a.unit, a.reason, a.by)
    print(f"{a.dst} 저장: {info['old']} → {info['new']} (행 {info['row']})")



ACK_COLS = ["일시", "실험ID", "검사", "확인 당시 판정값", "확인 당시 입력값", "확인 내용", "확인자"]


def acknowledge_check(src, dst, exp, check, note, by, when=None, sheet="실험기록"):
    """'확인 필요' 검사 중 계산설정 ACK_ALLOWED에 있는 검사에 대해, 연구자가 원자료와 대조해
    값이 맞다고 확인했음을 '검토확인' 시트에 남긴다. 확인 당시 원단위 계산에 쓰인 입력값
    (구간별 전력, 수소 생산량)을 함께 저장하며, 재실행 때 그중 하나라도 다르면 확인 표시는 무효가 된다.
    원본 파일은 바꾸지 않는다."""
    import carbon_receipt as cr
    wb_v = load_workbook(src, data_only=True)
    S = cr.Settings(wb_v)
    allowed = [x.strip() for x in str(S.val("ACK_ALLOWED")).split(",")]
    if check not in allowed:
        raise ValueError(f"'{check}'은 연구자 확인으로 넘어갈 수 없는 검사입니다(허용: {', '.join(allowed)})")
    EF = cr.load_ef(wb_v, str(S.val("EF_VERSION")).strip())
    recs = [r for r in cr.load_records(wb_v, sheet) if r["exp"] == exp]
    res = cr.evaluate(exp, recs, S, EF, [], [])
    st, detail = res["checks"][check]
    if st != cr.CHECK:
        raise ValueError(f"'{check}' 상태가 '{st}'라서 확인 표시를 남길 대상이 아닙니다")
    wb = load_workbook(src)
    if "검토확인" in wb.sheetnames:
        hs = wb["검토확인"]
    else:
        hs = wb.create_sheet("검토확인")
        for i, c in enumerate(ACK_COLS, start=1):
            cell = hs.cell(1, i, c)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            hs.column_dimensions[cell.column_letter].width = 50 if c in ("확인 내용", "확인 당시 입력값") else 18
    when = when or datetime.now().strftime("%Y-%m-%d %H:%M")
    hs.append([when, exp, check, detail, res["ack_keys"].get(check, detail), note, by])
    wb.save(dst)
    return detail


if __name__ == "__main__":
    main()
