"""검증 6사례 자동 확인 + 수작업 비교 + 오류 탐지 평가

실행: python run_validation.py
결과: 결과/ 폴더 (표는 CSV와 Markdown 두 형식, 사례별 영수증은 결과/영수증/)
모든 사례가 예상대로 처리되면 종료코드 0, 하나라도 다르면 1.
"""
import csv
import json
import math
import os
import shutil
import sys

from openpyxl import load_workbook

import carbon_receipt as cr
from correct_record import acknowledge_check, apply_correction
from make_cases import CASES, ELEC_CASE8, EXTRA_CASES, FORM, OUT_DIR as CASE_DIR, build as build_cases

RES = "결과"
RCPT = os.path.join(RES, "영수증")
HAND = "수작업기준.xlsx"


def run_case(path, name):
    out = os.path.join(RCPT, name)
    shutil.rmtree(out, ignore_errors=True)
    res = cr.run(path, out)
    assert len(res) == 1, f"{name}: 실험이 1건이어야 함"
    return res[0], out


def st(res, check):
    return res["checks"][check][0]


def write_table(stem, header, rows):
    with open(os.path.join(RES, stem + ".csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    md = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    md += ["| " + " | ".join(str(c).replace("|", "/").replace("\n", " ") for c in r) + " |" for r in rows]
    with open(os.path.join(RES, stem + ".md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md) + "\n")
    return "\n".join(md)


def hand_values():
    ws = load_workbook(HAND, data_only=True)["수작업계산"]
    vals = {}
    for r in range(1, ws.max_row + 1):
        a, e = ws.cell(r, 1).value, ws.cell(r, 5).value
        if a and isinstance(e, (int, float)):
            vals[a] = float(e)
    if "합계 — 배치 전체 · 부산물 공제 없음" not in vals:
        sys.exit("수작업기준.xlsx에 계산값이 저장되어 있지 않습니다. Excel에서 열어 저장한 뒤 다시 실행하세요.")
    return vals


def summarize(res):
    s = f"{res['state']}"
    flags = [f"{k} '{v[0]}'" for k, v in res["checks"].items() if v[0] != cr.PASS]
    if flags:
        s += " · " + ", ".join(flags)
    if res["total"] is not None:
        s += f" · 합계 {res['total']:.2f}"
    return s


def main():
    os.makedirs(RES, exist_ok=True)
    os.makedirs(RCPT, exist_ok=True)
    paths = build_cases()
    tol = cr.Settings(load_workbook(FORM, data_only=True)).num("CALC_TOL")
    H = hand_values()

    runs = {}
    for c in CASES:
        runs[c["key"]] = run_case(paths[c["key"]], c["key"])

    # ② 연구자 수정 모사: 원본 로그 대조 결과 Wh였다고 확인 → 수정본 저장 + 수정이력 기록 → 재검사·재계산
    fixed_path = os.path.join(CASE_DIR, "사례2_단위오기_수정본.xlsx")
    apply_correction(paths["사례2_단위오기"], fixed_path, "EXP-800", "반응", "ELEC", 5000, "Wh",
                     "원본 장비 로그 대조: 단위 오기(kWh→Wh)", "연구자(테스트 모사)", when="2026-10-09 10:00")
    runs["사례2_수정후"] = run_case(fixed_path, "사례2_단위오기_수정후")

    # ① 재실행 결과 동일 여부 (재현성 보조 확인)
    rerun, _ = run_case(paths["사례1_정상"], "사례1_정상_재실행")

    r1 = runs["사례1_정상"][0]
    # ---------------- 수작업 비교
    comp_spec = [
        ("전력 원단위 — 배치 전체", "kWh/kg-H2", r1["intermediate"]["전력 원단위-배치 전체 (kWh/kg-H2)"]),
        ("촉매 1회 배분량", "kg", r1["intermediate"]["촉매 1회 배분량 (kg)"]),
        ("전력(계통) 배출 — 배치 전체", "kgCO2-eq/kg-H2", r1["items"]["전력(계통)"]),
        ("외부 공급열 배출 — 배치 전체", "kgCO2-eq/kg-H2", r1["items"]["외부 공급열"]),
        ("원료 메탄 생산·공급 배출", "kgCO2-eq/kg-H2", r1["items"]["원료 메탄 생산·공급"]),
        ("촉매(Fe) 배출", "kgCO2-eq/kg-H2", r1["items"]["촉매(Fe)"]),
        ("합계 — 배치 전체 · 부산물 공제 없음", "kgCO2-eq/kg-H2", r1["total"]),
        ("합계 — 반응 구간만 · 부산물 공제 없음", "kgCO2-eq/kg-H2", r1["ref_total"]),
        ("배치 전체 · 질량 배분", "kgCO2-eq/kg-H2", r1["alloc"]),
    ]
    comp_rows, comp_ok, rels = [], True, []
    for name, unit, auto in comp_spec:
        hand = H[name]
        if hand == 0:
            rel = 0.0 if auto == 0 else math.inf
        else:
            rel = abs(auto - hand) / abs(hand)
        ok = rel <= tol
        comp_ok &= ok
        rels.append(rel)
        comp_rows.append([name, unit, f"{hand:.6f}", f"{auto:.6f}",
                          "0 (둘 다 0)" if hand == 0 and auto == 0 else f"{rel:.2e}", "일치" if ok else "불일치"])
    comp_md = write_table("수작업비교", ["항목", "단위", "수작업(Excel 수식)", "자동 계산", "상대오차",
                                       f"판정(허용 {tol:.1%})"], comp_rows)

    # ---------------- 사례별 예상 vs 실제
    rows, all_ok = [], True
    for c in CASES:
        res, out = runs[c["key"]]
        k = c["key"]
        png = os.path.exists(os.path.join(out, f"{res['exp']}_탄소영수증.png"))
        notes = []
        if k == "사례1_정상":
            ok = all(v[0] == cr.PASS for v in res["checks"].values()) and res["state"] == cr.STATE_DONE \
                and png and comp_ok
            actual = summarize(res) + f" · 영수증 출력 · 수작업 대비 최대 상대오차 {max(rels):.1e}"
        elif k == "사례2_단위오기":
            fx = runs["사례2_수정후"][0]
            same = all(abs(fx[f] - r1[f]) <= tol * abs(r1[f]) for f in ("total", "ref_total", "alloc"))
            ok = (st(res, c["target"]) == c["target_status"] and res["state"] == cr.STATE_HOLD
                  and fx["state"] == cr.STATE_DONE and all(v[0] == cr.PASS for v in fx["checks"].values())
                  and same and len(fx["history"]) == 1)
            actual = (summarize(res) + f" → 수정(5,000 kWh→5,000 Wh) 후 재검사: {fx['state']}, 합계 {fx['total']:.2f}"
                      f" (①과 {'동일' if same else '다름'}), 수정 이력 {len(fx['history'])}건 기록")
        elif k == "사례3_필수값누락":
            guide = [m for m in res["messages"] if "수소 생산량" in m and "입력해야" in m]
            ok = st(res, c["target"]) == c["target_status"] and res["state"] == cr.STATE_HOLD and bool(guide)
            actual = summarize(res) + (" · 안내: '수소 생산량(H2_PROD) 값을 입력해야 계산할 수 있음 — 측정 구간 …'"
                                       if guide else " · 안내 문구 없음")
        elif k == "사례4_중복기록":
            ws = load_workbook(paths[k], data_only=True)["실험기록"]
            n = sum(1 for r in range(2, ws.max_row + 1)
                    if ws.cell(r, 2).value == "반응" and ws.cell(r, 4).value == "ELEC")
            ok = st(res, c["target"]) == c["target_status"] and res["state"] == cr.STATE_HOLD and n == 2
            actual = summarize(res) + f" · 처리 후 입력 파일의 반응·전력 행 {n}개 유지(자동 삭제 안 함)"
        elif k == "사례5_수지불일치":
            causes = any(all(w in m for w in ("입력 오류", "반응기 내 축적", "미측정 유출")) for m in res["messages"])
            ok = (st(res, c["target"]) == c["target_status"] and res["state"] == cr.STATE_DONE
                  and res["total"] is not None and causes and png)
            actual = summarize(res) + (" · 영수증에 경고와 가능 원인(입력 오류/반응기 내 축적/미측정 유출) 표시"
                                       if causes else " · 가능 원인 문구 없음")
        elif k == "사례6_수소0":
            ok = (st(res, c["target"]) == c["target_status"] and res["state"] == cr.STATE_IMPOSSIBLE
                  and res["total"] is None and "0 이하" in res["state_reason"])
            actual = summarize(res) + f" · 사유: {res['state_reason']} (예외 없이 종료)"
        all_ok &= ok
        rows.append([c["no"] + " " + c["title"], c["inject"], c["expect"], actual, "일치" if ok else "불일치"])
    same_rerun = json.dumps([rerun["total"], rerun["ref_total"], rerun["alloc"]]) == \
        json.dumps([r1["total"], r1["ref_total"], r1["alloc"]])
    all_ok &= same_rerun
    case_md = write_table("6사례_결과표", ["사례", "입력", "예상 처리", "실제 처리(자동 확인)", "판정"], rows)

    # ---------------- 탐지율·오경보율
    det_rows, injected, detected = [], 0, 0
    fa_items, fa_count = 0, 0
    eval_runs = [(c, runs[c["key"]][0]) for c in CASES] + \
                [(dict(no="②", title="단위 오기 수정 후", target=None, related=[]), runs["사례2_수정후"][0])]
    for c, res in eval_runs:
        normal = [n for n in cr.ORIGINAL_CHECKS if n not in c["related"]]
        alarms = [n for n in normal if res["checks"][n][0] != cr.PASS]
        fa_items += len(normal)
        fa_count += len(alarms)
        if c.get("target"):
            injected += 1
            hit = res["checks"][c["target"]][0] == c["target_status"]
            detected += hit
            det = f"{'탐지' if hit else '미탐지'} ({c['target']}: {res['checks'][c['target']][0]})"
        else:
            det = "주입 오류 없음"
        det_rows.append([f"{c['no']} {c['title']}", det, len(normal), len(alarms), ", ".join(alarms) or "-"])
    det_rows.append(["합계", f"탐지율 {detected}/{injected} = {detected / injected:.0%}", fa_items, fa_count,
                     f"오경보율 {fa_count}/{fa_items} = {fa_count / fa_items:.1%}"])
    det_md = write_table("탐지평가", ["사례", "의도 오류 탐지", "정상 검사 항목 수", "정상 항목 중 경고 수", "비고"],
                         det_rows)
    with open(os.path.join(RES, "탐지평가_정의.md"), "w", encoding="utf-8") as f:
        f.write(
            "- 탐지율 = 탐지한 의도 오류 수 ÷ 주입한 오류 수. 사례 ②~⑥에 오류를 1개씩 주입하고, 해당 오류를 맡은 검사가\n"
            "  예상 상태(확인 필요/자료 부족/계산 불가)를 표시하면 탐지로 본다.\n"
            "- 오경보율 = 정상 항목 중 경고가 뜬 비율. 검사 항목(8개) 단위로 세며, 주입 오류를 맡은 검사와 주입으로 바뀐\n"
            "  입력값을 사용하는 검사(예: 수소 생산량을 비우면 전력 원단위·화학양론(H) 점검도 판정 불가)는 '관련 항목'으로\n"
            "  제외한다. 나머지 항목이 '통과'가 아니면 오경보로 센다. 사례 ①과 ② 수정 후 재검사는 8개 항목이 모두 정상 항목이다.\n"
            "- 사례 수가 6개로 적어 이 비율은 검사 규칙이 의도대로 동작하는지 확인하는 용도이며, 실제 자료에서의 성능을 뜻하지 않는다.\n")

    # 판정표: 7회 실행 × 8개 검사 = 56칸, 주입 오류 관련 칸 표시 (리뷰 A4)
    grid = []
    for c, res in eval_runs:
        for n in cr.ORIGINAL_CHECKS:
            stt, det = res["checks"][n]
            rel = n in c["related"]
            grid.append([f"{c['no']} {c['title']}", n, stt, "관련(제외)" if rel else "정상 항목",
                         "" if rel else ("경고" if stt != cr.PASS else "-"), det])
    write_table("탐지평가_판정표", ["실행", "검사", "상태", "구분", "오경보 여부", "세부"], grid)
    n_rel = sum(1 for g in grid if g[3] == "관련(제외)")
    with open(os.path.join(RES, "탐지평가_정의.md"), "a", encoding="utf-8") as f:
        f.write(f"- 판정표(탐지평가_판정표.csv): 6개 사례와 ② 수정 후 재검사, 모두 {len(eval_runs)}회 실행 × 8개 검사 = "
                f"{len(grid)}칸(처음 정한 8개 검사 기준). 이 중 주입 오류와 관련된 {n_rel}칸을 제외한 {len(grid) - n_rel}칸이 오경보율의 분모다. "
                "3차 검토 후 추가한 '값 범위·물리 한계' 검사는 이 표에 넣지 않고, 같은 7회 실행에서 모두 통과했는지만 따로 확인한다.\n"
                "- 관련 칸 기준(사례별): ② 전력 원단위 이상치 / ③ 필수값·수소 생산량·전력 원단위 이상치·화학양론(H) / "
                "④ 중복 기록·전력 원단위 이상치 / ⑤ 화학양론(C) / ⑥ 수소 생산량·전력 원단위 이상치·화학양론(H)\n")

    # ---------------- 추가 검증 ⑦⑧ (리뷰 B3) — 탐지 평가와 별도
    ext_rows, ext_ok = [], True
    r7, _ = run_case(paths["사례7_단위표기"], "사례7_단위표기")
    diffs = [abs(r7[f] - r1[f]) / abs(r1[f]) for f in ("total", "ref_total", "alloc")]
    ok7 = all(v[0] == cr.PASS for v in r7["checks"].values()) and r7["state"] == cr.STATE_DONE and max(diffs) <= tol
    ext_rows.append([EXTRA_CASES[0]["no"] + " " + EXTRA_CASES[0]["title"], EXTRA_CASES[0]["inject"],
                     EXTRA_CASES[0]["expect"],
                     f"{summarize(r7)} · ①과 최대 상대오차 {max(diffs):.1e}", "일치" if ok7 else "불일치"])
    ext_ok &= ok7

    p8 = paths["사례8_범위밖정상"]
    r8a, _ = run_case(p8, "사례8_범위밖정상_확인전")
    ack_path = os.path.join(CASE_DIR, "사례8_범위밖정상_확인표시.xlsx")
    acknowledge_check(p8, ack_path, "EXP-800", "전력 원단위 이상치",
                      "원본 장비 로그와 대조, 값이 맞음(테스트 모사)", "연구자(테스트 모사)", when="2026-10-09 11:00")
    r8b, _ = run_case(ack_path, "사례8_범위밖정상_확인후")
    # 독립 기대값: 입력에서 직접 계산 (전력 kWh ÷ H2 kg × 계수 + 메탄 + 촉매, 계수는 양식에서 읽음)
    ef = cr.load_ef(load_workbook(FORM, data_only=True), "EF-DEMO-v0.1")
    kwh8 = sum(ELEC_CASE8.values()) / 1000
    exp8 = kwh8 / 0.25 * ef["ELEC"]["value"] + 1.0 / 0.25 * ef["CH4_NEW"]["value"] + 0.02 / 4 / 0.25 * ef["CAT_MASS:Fe"]["value"]
    # 확인 후 값이 바뀌면 확인 표시가 무효가 되는지
    changed = os.path.join(CASE_DIR, "사례8_범위밖정상_확인후값변경.xlsx")
    apply_correction(ack_path, changed, "EXP-800", "반응", "ELEC", 26000, "Wh", "테스트: 확인 이후 값 변경",
                     "테스트", when="2026-10-09 11:10")
    r8c, _ = run_case(changed, "사례8_범위밖정상_값변경")
    # ⑧-2: 원단위는 같지만 입력이 바뀐 경우(승온 +1,000 Wh, 반응 −1,000 Wh) — 2차 검토에서 찾은 허점
    shift1 = os.path.join(CASE_DIR, "사례8-2_구간이동_1.xlsx")
    shift2 = os.path.join(CASE_DIR, "사례8-2_구간이동.xlsx")
    apply_correction(ack_path, shift1, "EXP-800", "승온", "ELEC", ELEC_CASE8["승온"] + 1000, "Wh",
                     "테스트: 승온 +1,000 Wh", "테스트", when="2026-10-09 11:20")
    apply_correction(shift1, shift2, "EXP-800", "반응", "ELEC", ELEC_CASE8["반응"] - 1000, "Wh",
                     "테스트: 반응 −1,000 Wh", "테스트", when="2026-10-09 11:21")
    os.remove(shift1)
    r82, _ = run_case(shift2, "사례8-2_구간이동")
    same_intensity = r82["checks"]["전력 원단위 이상치"][1] == r8a["checks"]["전력 원단위 이상치"][1]
    ok82 = (same_intensity and r82["checks"]["전력 원단위 이상치"][0] == cr.CHECK
            and r82["state"] == cr.STATE_HOLD and any("확인 표시 무효" in m for m in r82["messages"]))
    blocked = True
    try:
        acknowledge_check(paths["사례4_중복기록"], os.path.join(CASE_DIR, "_x.xlsx"), "EXP-800", "중복 기록", "x", "x")
        blocked = False
    except ValueError:
        pass
    ok8 = (r8a["checks"]["전력 원단위 이상치"][0] == cr.CHECK and r8a["state"] == cr.STATE_HOLD
           and r8b["checks"]["전력 원단위 이상치"][0] == cr.ACKED and r8b["state"] == cr.STATE_DONE
           and abs(r8b["total"] - exp8) <= tol * exp8 and "경고 유지" in r8b["state_reason"]
           and r8c["checks"]["전력 원단위 이상치"][0] == cr.CHECK and r8c["state"] == cr.STATE_HOLD
           and any("확인 표시 무효" in m for m in r8c["messages"]) and blocked)
    ext_rows.append([EXTRA_CASES[1]["no"] + " " + EXTRA_CASES[1]["title"], EXTRA_CASES[1]["inject"],
                     EXTRA_CASES[1]["expect"],
                     f"확인 전: {r8a['state']}({r8a['checks']['전력 원단위 이상치'][1]}) → 확인 표시 후: {r8b['state']}, "
                     f"'{r8b['checks']['전력 원단위 이상치'][0]}', 합계 {r8b['total']:.2f}(기대값 {exp8:.2f}) → "
                     f"반응 전력을 26,000 Wh로 바꾸면 확인 표시 무효·다시 보류 · 중복 기록에는 확인 표시를 남길 수 없음"
                     f"({'거부됨' if blocked else '허용됨'})", "일치" if ok8 else "불일치"])
    ext_ok &= ok8
    ext_rows.append(["⑧-2 원단위는 같고 입력만 바뀐 경우", "⑧의 확인 표시 후 승온 전력 +1,000 Wh, 반응 전력 −1,000 Wh "
                     "(원단위 130 kWh/kg-H₂ 그대로)", "확인 표시 무효, 다시 보류",
                     f"원단위 {r82['checks']['전력 원단위 이상치'][1]}({'같음' if same_intensity else '다름'}) · "
                     f"'{r82['checks']['전력 원단위 이상치'][0]}' · {r82['state']} · 확인 표시 무효 안내 출력",
                     "일치" if ok82 else "불일치"])
    ext_ok &= ok82
    all_ok &= ext_ok
    ext_md = write_table("추가검증_결과표", ["사례", "입력", "예상 처리", "실제 처리(자동 확인)", "판정"], ext_rows)
    if os.path.exists(os.path.join(CASE_DIR, "_x.xlsx")):
        os.remove(os.path.join(CASE_DIR, "_x.xlsx"))

    # 3차 검토 후 추가한 검사가 기존 7회 실행에서 오경보를 내지 않는지
    new_checks = [n for n in cr.CHECK_NAMES if n not in cr.ORIGINAL_CHECKS]
    new_ok = all(res["checks"][n][0] == cr.PASS for _, res in eval_runs for n in new_checks)
    all_ok &= new_ok

    # ---------------- ⑨ 3차 외부 검토 반례(R1~R5)와 자체 반례 점검(P01~) — 모두 계산 완료로 통과하면 안 됨
    import csv as _csv
    import probe_cases as pc
    ppaths = pc.build()
    prow = pc.run(cr, ppaths, os.path.join(RCPT, "반례"))
    before = {}
    if os.path.exists("반례_수정전_기록.csv"):
        for r in _csv.DictReader(open("반례_수정전_기록.csv", encoding="utf-8-sig")):
            before[r["ID"]] = r["수정 전 결과(엔진 v0.2, 2026-10-09 3차 검토 직후)"]
    exp_txt = {"hold": "계산 보류", "impossible": "계산 불가", "formerror": "양식 오류로 중단", "warn_done": "경고 후 계산(설계상 허용)", "done": "계산 완료(정상 범위로 판단)"}
    r_rows = [[p["id"], p["desc"], before.get(p["id"], "-"), exp_txt[p["expect"]], p["summary"], "일치" if p["ok"] else "불일치"]
              for p in prow if p["id"].startswith("R")]
    p_rows = [[p["id"], p["desc"], before.get(p["id"], "-"), exp_txt[p["expect"]], p["summary"], "일치" if p["ok"] else "불일치"]
              for p in prow if p["id"].startswith("P")]
    hdr = ["ID", "변경", "수정 전 결과", "예상 처리", "수정 후 결과", "판정"]
    r_md = write_table("검토반례_결과표", hdr, r_rows)
    p_md = write_table("자체반례_점검표", hdr, p_rows)
    probe_ok = all(p["ok"] for p in prow)
    all_ok &= probe_ok

    # 정상 사례 영수증 대표 그림
    shutil.copy(os.path.join(RCPT, "사례1_정상", "EXP-800_탄소영수증.png"), os.path.join(RES, "정상사례_탄소영수증.png"))

    print("=== 6사례 예상 vs 실제 ===\n" + case_md)
    print("\n=== 수작업 vs 자동 계산 (사례 ①) ===\n" + comp_md)
    print("\n=== 오류 탐지 평가 ===\n" + det_md)
    print("\n=== 추가 검증 ⑦⑧⑧-2 (탐지 평가와 별도) ===\n" + ext_md)
    print("\n=== ⑨ 3차 외부 검토 반례 ===\n" + r_md)
    print("\n=== 자체 반례 점검 ===\n" + p_md)
    print(f"\n추가 검사({', '.join(new_checks)})가 기존 7회 실행에서 모두 통과: {'예' if new_ok else '아니오'}")
    print(f"\n재실행 결과 동일(사례 ①): {'예' if same_rerun else '아니오'}")
    print(f"\n전체 판정: {'모두 예상대로 처리됨' if all_ok else '예상과 다른 사례 있음'}")
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
