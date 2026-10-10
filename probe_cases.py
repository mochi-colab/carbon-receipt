"""비정상 입력 반례 점검

정상 사례(①) 입력 파일을 하나씩 비정상으로 바꾼 뒤 엔진을 실행해, '계산 완료'로 통과하는 경우가
없는지 확인한다. R1~R5는 3차 외부 검토에서 발견된 반례(⑨), P01~는 팀이 추가로 찔러본 경우이다.
expect: "hold" = 계산 보류, "impossible" = 계산 불가, "formerror" = 양식 오류로 실행 중단,
        "warn_done" = 경고를 남기고 계산(설계상 허용)
"""
import os
import shutil

from openpyxl import load_workbook

CASE1 = os.path.join("검증사례", "사례1_정상.xlsx")
OUT = os.path.join("검증사례", "반례")


def _rows(wb):
    ws = wb["실험기록"]
    hdr = {c.value: c.column for c in ws[1] if c.value}
    return ws, hdr


def set_val(seg, code, value, unit=None):
    def f(wb):
        ws, h = _rows(wb)
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, h["구간"]).value == seg and ws.cell(r, h["항목코드"]).value == code:
                ws.cell(r, h["값"]).value = value
                if unit is not None:
                    ws.cell(r, h["단위"]).value = unit
                return
        raise KeyError((seg, code))
    return f


def del_row(seg, code):
    def f(wb):
        ws, h = _rows(wb)
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, h["구간"]).value == seg and ws.cell(r, h["항목코드"]).value == code:
                ws.delete_rows(r)
                return
        raise KeyError((seg, code))
    return f


def add_row(seg, name, code, value, unit):
    def f(wb):
        ws, h = _rows(wb)
        r = ws.max_row + 1
        while r > 2 and not ws.cell(r - 1, h["실험ID"]).value:
            r -= 1
        for k, v in (("실험ID", "EXP-800"), ("구간", seg), ("항목", name), ("항목코드", code), ("값", value),
                     ("단위", unit), ("필수/선택", "선택")):
            ws.cell(r, h[k]).value = v
    return f


def set_ef(ef_id, col, value):
    def f(wb):
        ws = wb["배출계수"]
        h = {c.value: c.column for c in ws[1] if c.value}
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == ef_id:
                ws.cell(r, h[col]).value = value
                return
        raise KeyError(ef_id)
    return f


def dup_ef(ef_id, value):
    def f(wb):
        ws = wb["배출계수"]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == ef_id:
                vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                vals[0] = ef_id + "-B"
                vals[2] = value
                ws.append(vals)
                return
        raise KeyError(ef_id)
    return f


def set_setting(sid, value):
    def f(wb):
        ws = wb["계산설정"]
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == sid:
                ws.cell(r, 4).value = value
                return
        raise KeyError(sid)
    return f


# (ID, 설명, 변경 함수, 예상)
PROBES = [
    ("R1", "메탄 신규 공급량 −1 kg", set_val("반응", "CH4_NEW", -1), "hold"),
    ("R2", "반응온도 행 삭제", del_row("반응", "TEMP"), "hold"),
    ("R3", "전력 배출계수 단위를 kgCO2-eq/MWh로 변경(값 0.45 그대로)", set_ef("EF-ELEC-01", "단위", "kgCO2-eq/MWh"), "hold"),
    ("R4", "촉매 사용량 −20 g", set_val("반응", "CAT_MASS", -20), "hold"),
    ("R5", "반응 구간 외부 공급열 −5 MJ", set_val("반응", "HEAT_EXT", -5), "hold"),
    ("P01", "승온 전력 −1,200 Wh", set_val("승온", "ELEC", -1200), "hold"),
    ("P02", "반응 운전시간 −4 h", set_val("반응", "OP_TIME", -4), "hold"),
    ("P03", "고체탄소 회수량 −720 g", set_val("반응", "C_REC", -720), "hold"),
    ("P04", "촉매 총 사용 횟수 0", set_val("반응", "CAT_USES", 0), "hold"),
    ("P05", "촉매 총 사용 횟수 2.5(정수 아님)", set_val("반응", "CAT_USES", 2.5), "hold"),
    ("P06", "출구 H2 몰분율 1.5", set_val("반응", "X_H2", 1.5), "hold"),
    ("P07", "수소 순도 150 mol%", set_val("반응", "H2_PURITY", 150), "hold"),
    ("P08", "반응온도 −800 ℃", set_val("반응", "TEMP", -800), "hold"),
    ("P09", "미반응 메탄 배출량 3 kg(투입 2 kg보다 큼)", set_val("반응", "CH4_OUT", 3), "hold"),
    ("P10", "메탄 신규 공급량 0 kg(수소 0.25 kg 생산)", set_val("반응", "CH4_NEW", 0), "hold"),
    ("P11", "메탄 신규 공급량 0.5 kg(수소 0.25 kg에 필요한 이론 최소량 약 0.99 kg 미만)",
     set_val("반응", "CH4_NEW", 0.5), "hold"),
    ("P12", "등록되지 않은 구간 '냉각'에 전력 1,000 Wh 추가", add_row("냉각", "전력 사용량", "ELEC", 1000, "Wh"), "hold"),
    ("P13", "등록되지 않은 항목코드 ELEC_AUX 1,000 Wh 추가", add_row("반응", "보조 전력", "ELEC_AUX", 1000, "Wh"), "hold"),
    ("P14", "전력 배출계수 값 −0.45", set_ef("EF-ELEC-01", "값", -0.45), "hold"),
    ("P15", "같은 버전에 전력 배출계수 2개(0.45, 0.10)", dup_ef("EF-ELEC-01", 0.10), "hold"),
    ("P16", "메탄 배출계수 단위를 kgCO2-eq/m3로 변경", set_ef("EF-CH4-01", "단위", "kgCO2-eq/m3"), "hold"),
    ("P17", "수소 생산량 행 삭제", del_row("반응", "H2_PROD"), "hold"),
    ("P18", "수소 생산량 −0.25 kg", set_val("반응", "H2_PROD", -0.25, "kg"), "impossible"),
    ("P19", "전력 단위 소문자 'wh'", set_val("반응", "ELEC", 5000, "wh"), "hold"),
    ("P20", "계산설정 기능단위(FU) 0", set_setting("FU", 0), "formerror"),
    ("P21", "출구 CH4 몰분율 −0.1", set_val("반응", "X_CH4", -0.1), "hold"),
    # 2차 점검(규칙 추가 후)
    ("P22", "촉매 종류 Ni(해당 배출계수 없음)", set_val("반응", "CAT_TYPE", "Ni"), "hold"),
    ("P23", "반응 전력 행의 실험ID 공란", lambda wb: _blank_exp(wb, "반응", "ELEC"), "hold"),
    ("P24", "반응 전력 값에 문자 'abc'", set_val("반응", "ELEC", "abc"), "hold"),
    ("P25", "계산설정의 배출계수 버전을 없는 버전으로 변경", set_setting("EF_VERSION", "EF-DEMO-v9"), "hold"),
    ("P26", "수소 생산량 1,000 kg", set_val("반응", "H2_PROD", 1000, "kg"), "hold"),
    ("P27", "전력 3구간 모두 0", lambda wb: [set_val(sg, "ELEC", 0)(wb) for sg in ("승온", "반응", "중단")], "hold"),
    ("P28", "외부 공급열 단위를 kWh로 입력", set_val("반응", "HEAT_EXT", 0, "kWh"), "hold"),
    ("P29", "출구 몰분율 합 1.3(CH4 0.7, H2 0.6)",
     lambda wb: [set_val("반응", "X_CH4", 0.7)(wb), set_val("반응", "X_H2", 0.6)(wb)], "hold"),
    ("P30", "반응 운전시간 0 h(수소는 생산)", set_val("반응", "OP_TIME", 0), "hold"),
    ("P31", "고체탄소 회수량 2 kg(이론량의 약 2.7배)", set_val("반응", "C_REC", 2000), "warn_done"),
    ("P32", "수소 생산량 0.3 kg(메탄 1 kg 기준 이론량의 119%)", set_val("반응", "H2_PROD", 300), "hold"),
    ("P33", "메탄 신규 공급량 100 kg(반응량보다 훨씬 많음)", set_val("반응", "CH4_NEW", 100), "done"),
    ("P34", "전력 배출계수 값을 문자 'nan'으로 입력", set_ef("EF-ELEC-01", "값", "nan"), "hold"),
    ("P35", "전력 배출계수 값을 문자 'inf'로 입력", set_ef("EF-ELEC-01", "값", "inf"), "hold"),
]


def _blank_exp(wb, seg, code):
    ws, h = _rows(wb)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, h["구간"]).value == seg and ws.cell(r, h["항목코드"]).value == code:
            ws.cell(r, h["실험ID"]).value = None
            return
    raise KeyError((seg, code))


def build():
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT)
    paths = {}
    for pid, desc, fn, exp in PROBES:
        wb = load_workbook(CASE1)
        fn(wb)
        p = os.path.join(OUT, f"{pid}.xlsx")
        wb.save(p)
        paths[pid] = p
    return paths


def run(engine, paths, out_root):
    """engine = carbon_receipt 모듈(수정 전/후 비교용으로 다른 모듈을 넘길 수 있음)"""
    rows = []
    for pid, desc, fn, exp in PROBES:
        try:
            res = engine.run(paths[pid], os.path.join(out_root, pid), png=False)[0]
            state = res["state"]
            flags = [f"{k} '{s}'" for k, (s, d) in res["checks"].items() if s != engine.PASS]
            tot = f"{res['total']:.2f}" if res["total"] is not None else "-"
            got = {"계산 완료": "done", "계산 보류": "hold", "계산 불가": "impossible"}[state]
            if got == "done" and flags:
                got = "warn_done"
            summary = f"{state}" + (f" {tot}" if res["total"] is not None else "") + \
                      (f" · {', '.join(flags)}" if flags else " · 모든 검사 통과")
        except engine.FormError as e:
            got, summary = "formerror", f"양식 오류로 중단: {e}"
        except Exception as e:  # 프로그램 오류(예외)로 멈춘 경우도 기록
            got, summary = "crash", f"프로그램 오류: {type(e).__name__}: {e}"
        rows.append(dict(id=pid, desc=desc, expect=exp, got=got, summary=summary, ok=(got == exp)))
    return rows
