"""수작업기준.xlsx 생성 — 사례 ①(작성예시 EXP-800)을 Excel 수식으로 직접 계산

계산 엔진과 독립적으로 검산하기 위해, 입력값과 배출계수를 양식에서 읽지 않고 작성예시를 보고 손으로 옮겨
적은 값(파란 글씨)으로 둔다. 모든 결과 칸은 Excel 수식이다.
(이 파일은 한 번 만들어 Excel/LibreOffice로 재계산해 두며, run_validation.py는 저장된 계산값을 읽는다.)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = "수작업기준.xlsx"
F = lambda **k: Font(**{"name": "Arial", "size": 10, **k})
BLUE, BLACK = "0000FF", "000000"
HDR = PatternFill("solid", fgColor="1F3864")
SEC = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INPUTS = [  # 키, 항목, 값, 입력단위, 환산계수, 표준단위
    ("e_heat", "승온 구간 전력 사용량", 1200, "Wh", 0.001, "kWh"),
    ("e_react", "반응 구간 전력 사용량", 5000, "Wh", 0.001, "kWh"),
    ("e_stop", "중단 구간 전력 사용량", 300, "Wh", 0.001, "kWh"),
    ("q_heat", "승온 구간 외부 공급열", 0, "MJ", 1, "MJ"),
    ("q_react", "반응 구간 외부 공급열", 0, "MJ", 1, "MJ"),
    ("q_stop", "중단 구간 외부 공급열", 0, "MJ", 1, "MJ"),
    ("ch4", "메탄 신규 공급량", 1.0, "kg", 1, "kg"),
    ("h2", "수소 생산량", 250, "g", 0.001, "kg"),
    ("crec", "고체탄소 회수량", 720, "g", 0.001, "kg"),
    ("cat", "촉매 사용량", 20, "g", 0.001, "kg"),
    ("uses", "촉매 최초 사용 포함 총 사용 횟수", 4, "회", 1, "회"),
]
EFS = [  # 키, 대상, 값, 단위
    ("ef_e", "전력(계통) — 예시값", 0.45, "kgCO2-eq/kWh"),
    ("ef_q", "외부 공급열 — 예시값", 0.07, "kgCO2-eq/MJ"),
    ("ef_m", "메탄 생산·공급 — 예시값", 0.6, "kgCO2-eq/kg-CH4"),
    ("ef_c", "촉매(Fe) — 예시값", 2.0, "kgCO2-eq/kg-촉매"),
]
# 키, 항목, 수식(키 치환), 단위, A 단계 검산값(사용자 확인값, 반올림 표기)
CALCS = [
    ("kwh_b", "전력 합계 — 배치 전체", "={e_heat}+{e_react}+{e_stop}", "kWh", None),
    ("kwh_r", "전력 합계 — 반응 구간", "={e_react}", "kWh", None),
    ("mj_b", "외부 공급열 합계 — 배치 전체", "={q_heat}+{q_react}+{q_stop}", "MJ", None),
    ("mj_r", "외부 공급열 합계 — 반응 구간", "={q_react}", "MJ", None),
    ("cat_a", "촉매 1회 배분량", "={cat}/{uses}", "kg", None),
    ("inten", "전력 원단위 — 배치 전체", "={kwh_b}/{h2}", "kWh/kg-H2", 26),
    ("i_e", "전력(계통) 배출 — 배치 전체", "={kwh_b}*{ef_e}/{h2}", "kgCO2-eq/kg-H2", 11.70),
    ("i_q", "외부 공급열 배출 — 배치 전체", "={mj_b}*{ef_q}/{h2}", "kgCO2-eq/kg-H2", 0),
    ("i_m", "원료 메탄 생산·공급 배출", "={ch4}*{ef_m}/{h2}", "kgCO2-eq/kg-H2", 2.40),
    ("i_c", "촉매(Fe) 배출", "={cat_a}*{ef_c}/{h2}", "kgCO2-eq/kg-H2", 0.04),
    ("tot_b", "합계 — 배치 전체 · 부산물 공제 없음", "={i_e}+{i_q}+{i_m}+{i_c}", "kgCO2-eq/kg-H2", 14.14),
    ("tot_r", "합계 — 반응 구간만 · 부산물 공제 없음",
     "={kwh_r}*{ef_e}/{h2}+{mj_r}*{ef_q}/{h2}+{i_m}+{i_c}", "kgCO2-eq/kg-H2", 11.44),
    ("share", "수소 질량 배분율", "={h2}/({h2}+{crec})", "-", None),
    ("alloc", "배치 전체 · 질량 배분", "={tot_b}*{share}", "kgCO2-eq/kg-H2", 3.644),
]


def build(path=OUT):
    wb = Workbook()
    ws = wb.active
    ws.title = "수작업계산"
    for col, w in zip("ABCDEFG", [38, 12, 12, 12, 16, 18, 18]):
        ws.column_dimensions[col].width = w
    ws["A1"] = "수작업 기준 계산 — 사례 ① (가상 실험 EXP-800, 예시 배출계수)"
    ws["A1"].font = F(bold=True, size=12)
    ws["A2"] = "파란 글씨 = 작성예시에서 손으로 옮긴 입력값 · 검은 글씨 = Excel 수식. 가상 데이터·예시 배출계수 기반."
    ws["A2"].font = F(color="C00000")
    ref = {}
    r = 4

    def head(cols):
        nonlocal r
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(r, i, c)
            cell.font = F(bold=True, color="FFFFFF"); cell.fill = HDR; cell.border = BOX
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        r += 1

    def section(t):
        nonlocal r
        ws.cell(r, 1, t).font = F(bold=True)
        for c in range(1, 8):
            ws.cell(r, c).fill = SEC
        r += 1

    section("1. 입력값 (작성예시 전사)")
    head(["항목", "입력값", "입력단위", "환산계수", "표준단위 값", "표준단위", ""])
    for k, name, v, u, fct, su in INPUTS:
        ws.cell(r, 1, name); ws.cell(r, 2, v).font = F(color=BLUE); ws.cell(r, 3, u)
        ws.cell(r, 4, fct).font = F(color=BLUE); ws.cell(r, 5, f"=B{r}*D{r}"); ws.cell(r, 6, su)
        ref[k] = f"$E${r}"
        r += 1
    r += 1
    section("2. 배출계수 (시연용 예시값, 공식 계수 아님)")
    head(["대상", "값", "단위", "", "", "", ""])
    for k, name, v, u in EFS:
        ws.cell(r, 1, name); ws.cell(r, 2, v).font = F(color=BLUE); ws.cell(r, 3, u)
        ref[k] = f"$B${r}"
        r += 1
    r += 1
    section("3. 계산 (수식)")
    head(["항목", "", "", "", "수작업 계산값", "단위", "A 단계 검산값"])
    for k, name, frm, u, a in CALCS:
        ws.cell(r, 1, name)
        ref[k] = f"$E${r}"
        ws.cell(r, 5, frm.format(**ref))
        ws.cell(r, 5).number_format = "0.000000"
        ws.cell(r, 6, u)
        if a is not None:
            ws.cell(r, 7, a).font = F(color=BLUE)
        r += 1
    for row in ws.iter_rows(min_row=4, max_row=r):
        for c in row:
            if c.value is not None and c.font.color is None:
                c.font = F()
    ws.cell(r + 1, 1, "주: 'A 단계 검산값'은 입력양식 작성예시의 검산 결과(반올림 표기)이며, 비교는 수작업 계산값 기준으로 한다.").font = F(color="777777")
    wb.save(path)
    return {k: v for k, v in ref.items()}


if __name__ == "__main__":
    print(build())
